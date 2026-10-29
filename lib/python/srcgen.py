#!/usr/bin/env python3
"""软著「源程序鉴别材料」生成器。

规则（中国版权保护中心 · 一般交存）：
  - 源程序提交前、后各连续 30 页；不足 60 页则全部提交
  - 页眉标注软件名称 + 版本号，右上角标注页码
  - 每页不少于 50 行

实现：每页写满 50 行后插入显式分页符，行距固定，因此分页与渲染器无关、可复现。

输出：最后一个说明块之后的 `##DSH_JSON##` 行给出机器可读结果。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

LINES_PER_PAGE = 50
PAGES_PER_SIDE = 30
MAX_LINE_LEN = 92          # Consolas 9pt 在 17cm 正文宽度内约容 96 字符
MAX_FILE_BYTES = 1024 * 1024
LINE_SPACING_PT = 11.5     # 每页可容 ≈63 行，留安全余量
FONT_SIZE_PT = 9

SKIP_DIRS = {
    "node_modules", ".git", ".venv", "venv", "dist", "build", "__pycache__",
    ".pytest_cache", ".ruff_cache", ".uv-cache", ".idea", ".vscode", "coverage",
    "vendor", "target", ".next", ".cache", ".mypy_cache", ".tox", "site-packages",
}
SKIP_SUFFIX = {
    ".lock", ".log", ".png", ".jpg", ".jpeg", ".gif", ".svg", ".ico", ".pdf",
    ".zip", ".tar", ".gz", ".woff", ".woff2", ".ttf", ".mp4", ".mp3", ".map",
    ".pyc", ".so", ".dll", ".exe", ".bin", ".jar", ".class",
}
CODE_SUFFIX = {
    ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".vue", ".java", ".kt",
    ".go", ".rs", ".c", ".h", ".cpp", ".hpp", ".cs", ".php", ".rb", ".swift",
    ".sql", ".sh", ".bash", ".html", ".css", ".scss", ".less", ".yml", ".yaml",
    ".toml", ".ini", ".conf", ".xml", ".proto", ".lua", ".dart", ".scala", ".pl",
}
COMMENT_PREFIX = {
    ".py": "#", ".sh": "#", ".bash": "#", ".yml": "#", ".yaml": "#", ".toml": "#",
    ".ini": "#", ".conf": "#", ".pl": "#", ".rb": "#",
    ".sql": "--",
    ".html": "<!--", ".xml": "<!--", ".vue": "<!--",
}


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def skipped(path: Path, root: Path, excludes: tuple[str, ...]) -> bool:
    rel = path.relative_to(root)
    if any(part in SKIP_DIRS for part in rel.parts[:-1]):
        return True
    if path.name.startswith("."):
        return True
    low = path.name.lower()
    if any(low.endswith(s) for s in SKIP_SUFFIX):
        return True
    if path.suffix.lower() not in CODE_SUFFIX:
        return True
    posix = rel.as_posix()
    return any(ex and ex in posix for ex in excludes)


def collect(root: Path, excludes: tuple[str, ...]) -> tuple[list[str], list[str]]:
    files = sorted(p for p in root.rglob("*") if p.is_file() and not skipped(p, root, excludes))
    lines: list[str] = []
    used: list[str] = []
    for path in files:
        try:
            if path.stat().st_size > MAX_FILE_BYTES:
                continue
            text = read_text(path)
        except OSError:
            continue
        if not text.strip():
            continue
        rel = path.relative_to(root).as_posix()
        prefix = COMMENT_PREFIX.get(path.suffix.lower(), "//")
        used.append(rel)
        if prefix == "<!--":
            lines.append(f"<!-- ===== {rel} ===== -->")
        else:
            lines.append(f"{prefix} ===== {rel} =====")
        for raw in text.splitlines():
            line = raw.replace("\t", "    ").rstrip()
            if not line:
                continue
            while len(line) > MAX_LINE_LEN:
                lines.append(line[:MAX_LINE_LEN])
                line = "    " + line[MAX_LINE_LEN:]
            lines.append(line)
    return lines, used


def pick(lines: list[str]) -> tuple[list[str], str]:
    full = LINES_PER_PAGE * PAGES_PER_SIDE * 2
    if len(lines) <= full:
        return lines, "all"
    half = LINES_PER_PAGE * PAGES_PER_SIDE
    return lines[:half] + lines[-half:], "head-tail"


def header_run(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(FONT_SIZE_PT)
    run.font.name = "Consolas"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Consolas")
    rfonts.set(qn("w:eastAsia"), "宋体")


def page_field(paragraph) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Consolas")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(FONT_SIZE_PT * 2))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def build(lines: list[str], name: str, version: str, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)

    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.9), WD_TAB_ALIGNMENT.RIGHT)
    header_run(hp, f"{name}  {version}\t")
    page_field(hp)

    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(FONT_SIZE_PT)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_SPACING_PT)

    pages = (len(lines) + LINES_PER_PAGE - 1) // LINES_PER_PAGE
    for page in range(pages):
        for i in range(LINES_PER_PAGE):
            idx = page * LINES_PER_PAGE + i
            if idx >= len(lines):
                break
            p = doc.add_paragraph()
            pf2 = p.paragraph_format
            pf2.space_before = Pt(0)
            pf2.space_after = Pt(0)
            pf2.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf2.line_spacing = Pt(LINE_SPACING_PT)
            # 段前分页：比插入 break 运行更稳（各渲染器都认），且不占行
            if i == 0 and page > 0:
                pf2.page_break_before = True
            header_run(p, lines[idx])

    doc.core_properties.title = f"{name} {version} 源程序"
    doc.core_properties.subject = "计算机软件著作权登记 · 源程序鉴别材料"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--name", required=True)
    ap.add_argument("--version", default="V1.0")
    ap.add_argument("--out", required=True)
    ap.add_argument("--exclude", action="append", default=[])
    args = ap.parse_args()

    root = Path(args.root).expanduser().resolve()
    if not root.is_dir():
        print(f"root is not a directory: {root}", file=sys.stderr)
        return 2

    all_lines, files = collect(root, tuple(e for e in args.exclude if e))
    if not all_lines:
        print("no source files matched", file=sys.stderr)
        return 3

    lines, mode = pick(all_lines)
    pages = (len(lines) + LINES_PER_PAGE - 1) // LINES_PER_PAGE
    out = Path(args.out).expanduser().resolve()
    build(lines, args.name, args.version, out)

    print(f"files   : {len(files)}")
    print(f"lines   : {len(all_lines)} total -> {len(lines)} submitted ({mode})")
    available_pt = (29.7 - 2.0 * 2.0) / 2.54 * 72
    used_pt = LINES_PER_PAGE * LINE_SPACING_PT
    print(f"pages   : {pages} ({LINES_PER_PAGE} lines/page)")
    print(f"fit     : {used_pt:.0f}pt used / {available_pt:.0f}pt available "
          f"({'ok' if used_pt < available_pt else 'OVERFLOW'})")
    print(f"out     : {out}")
    print("##DSH_JSON## " + json.dumps({
        "ok": True,
        "docx": str(out),
        "software_name": args.name,
        "version": args.version,
        "files": len(files),
        "total_lines": len(all_lines),
        "submitted_lines": len(lines),
        "pages": pages,
        "mode": mode,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
