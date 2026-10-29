#!/usr/bin/env python3
"""软著《软件说明书》生成器（排版层）。

只负责排「合规的样子」，内容由 agent / 人提供：
  - A4，页边距 2.54 / 3.17cm，正文宋体小四、1.5 倍行距（每页 ≥30 行）
  - 页眉标注软件名称 + 版本号，右上角页码
  - 标题自动编号（1 / 1.1 / 1.1.1）
  - 可选目录域（Word 打开后按 F9 更新）
  - 图片居中嵌入 + 图注，表格套用 Table Grid

输入：UTF-8 JSON 文件
输出：最后一个说明块之后的 `##DSH_JSON##` 行给出机器可读结果。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_PT = 12          # 小四
H1_PT, H2_PT, H3_PT = 16, 14, 12
BODY_LINE = 1.5
MAX_IMG_WIDTH_CM = 14.0


def set_east_asian(style_or_run, font_name: str, size_pt: float | None = None, bold: bool | None = None) -> None:
    font = style_or_run.font
    font.name = font_name
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    el = style_or_run.element
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        el.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), "Times New Roman" if font_name == "宋体" else font_name)
    rfonts.set(qn("w:hAnsi"), "Times New Roman" if font_name == "宋体" else font_name)


def page_field(paragraph) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "宋体")
    rfonts.set(qn("w:hAnsi"), "宋体")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")  # 9pt
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "（在此处右键 → 更新域，生成目录）"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, text, end):
        run._r.append(node)


def setup(doc: Document, header_left: str) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    sec.header_distance = Cm(1.5)

    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(14.66), WD_TAB_ALIGNMENT.RIGHT)
    run = hp.add_run(f"{header_left}\t")
    set_east_asian(run, "宋体", 9)
    page_field(hp)

    normal = doc.styles["Normal"]
    set_east_asian(normal, "宋体", BODY_PT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(BODY_PT * 2)


class Stats:
    """按 Word 的排版密度估算页数，用于提醒「篇幅不足」。"""

    CHARS_PER_LINE = 33
    LINES_PER_PAGE = 30

    def __init__(self) -> None:
        self.chars = 0
        self.paragraphs = 0
        self.bullets = 0
        self.numbered = 0
        self.table_rows = 0
        self.table_headers = 0
        self.images = 0
        self.h1 = self.h2 = self.h3 = 0

    def lines(self) -> float:
        return (
            self.chars / self.CHARS_PER_LINE
            + self.paragraphs * 0.32
            + (self.bullets + self.numbered) * 0.8
            + self.table_rows * 1.6
            + self.table_headers * 1.2
            + self.images * 14
            + self.h1 * 2.2 + self.h2 * 1.8 + self.h3 * 1.6
        )

    def pages(self) -> int:
        # +1 给目录/封面留一页
        return max(1, round(self.lines() / self.LINES_PER_PAGE) + 1)


class Numberer:
    def __init__(self) -> None:
        self.counters = [0, 0, 0]

    def next(self, level: int) -> str:
        level = max(1, min(3, level))
        self.counters[level - 1] += 1
        for i in range(level, 3):
            self.counters[i] = 0
        return ".".join(str(c) for c in self.counters[:level])


def add_heading(doc, text: str, level: int, number: str, page_break_before: bool = False) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if page_break_before:
        pf.page_break_before = True
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(f"{number}  {text}")
        set_east_asian(run, "黑体", H1_PT, True)
    elif level == 2:
        run = p.add_run(f"{number}  {text}")
        set_east_asian(run, "黑体", H2_PT, True)
    else:
        run = p.add_run(f"{number}  {text}")
        set_east_asian(run, "黑体", H3_PT, True)


def add_body(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_east_asian(run, "宋体", BODY_PT)


def add_numbered(doc, index: int, text: str) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(BODY_PT * 2.4)
    pf.space_after = Pt(3)
    run = p.add_run(f"{index}. {text}")
    set_east_asian(run, "宋体", BODY_PT)


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(BODY_PT * 2)
    pf.space_after = Pt(3)
    run = p.add_run(f"· {text}")
    set_east_asian(run, "宋体", BODY_PT)


def add_image(doc, path: Path, caption: str, index: int) -> bool:
    if not path.is_file():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(MAX_IMG_WIDTH_CM))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.first_line_indent = Pt(0)
        c.paragraph_format.space_after = Pt(8)
        run = c.add_run(caption)
        set_east_asian(run, "宋体", 10.5)
    return True


def add_table(doc, spec: dict) -> int:
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    if not header and not rows:
        return 0
    cols = len(header) or max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = table.add_row().cells
        for i, text in enumerate(header[:cols]):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(text))
            set_east_asian(run, "黑体", 10.5, True)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row[:cols]):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(text))
            set_east_asian(run, "宋体", 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return len(rows) + (1 if header else 0)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True, help="UTF-8 JSON 说明文件")
    ap.add_argument("--out", required=True)
    ap.add_argument("--append", help="追加的章节 JSON（例如 codeextract.py 生成的附录）")
    args = ap.parse_args()

    spec_path = Path(args.spec).expanduser().resolve()
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"bad spec: {exc}", file=sys.stderr)
        return 2

    if args.append:
        try:
            extra = json.loads(Path(args.append).expanduser().read_text(encoding="utf-8"))
            spec.setdefault("sections", []).extend(extra.get("sections") or [])
        except (OSError, json.JSONDecodeError) as exc:
            print(f"附录读取失败：{exc}", file=sys.stderr)
            return 4

    name = spec.get("software_name") or "未命名软件"
    version = spec.get("version") or "V1.0"
    sections = spec.get("sections") or []
    if not sections:
        print("spec.sections is empty", file=sys.stderr)
        return 3

    doc = Document()
    setup(doc, f"{name}  {version}")

    if spec.get("toc", True):
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Pt(0)
        run = title.add_run("目  录")
        set_east_asian(run, "黑体", 16, True)
        tp = doc.add_paragraph()
        tp.paragraph_format.first_line_indent = Pt(0)
        toc_field(tp)

    stats = Stats()
    numberer = Numberer()
    first_heading = bool(spec.get("toc", True))
    images = 0
    image_index = 0
    tables = 0
    chars = 0
    for section in sections:
        heading = (section.get("heading") or "").strip()
        level = int(section.get("level") or 1)
        if heading:
            add_heading(doc, heading, level, numberer.next(level), page_break_before=first_heading)
            first_heading = False
            if level == 1:
                stats.h1 += 1
            elif level == 2:
                stats.h2 += 1
            else:
                stats.h3 += 1
        for para in section.get("paragraphs") or []:
            text = str(para).strip()
            if text:
                add_body(doc, text)
                chars += len(text)
                stats.chars += len(text)
                stats.paragraphs += 1
        for i, bullet in enumerate(section.get("bullets") or [], start=1):
            text = str(bullet).strip()
            if text:
                add_bullet(doc, text)
                chars += len(text)
                stats.chars += len(text)
                stats.bullets += 1

        for i, step in enumerate(section.get("steps") or [], start=1):
            text = str(step).strip()
            if text:
                add_numbered(doc, i, text)
                chars += len(text)
                stats.chars += len(text)
                stats.numbered += 1
        table_spec = section.get("table")
        if isinstance(table_spec, dict):
            rows_added = add_table(doc, table_spec)
            tables += rows_added
            stats.table_rows += len(table_spec.get("rows") or [])
            stats.table_headers += 1 if table_spec.get("header") else 0
        image = section.get("image")
        if image:
            image_index += 1
            caption = section.get("caption") or f"图 {image_index}  {section.get('heading', '')}"
            if add_image(doc, Path(image).expanduser(), caption, image_index):
                images += 1
                stats.images += 1

    doc.core_properties.title = f"{name} {version} 软件说明书"
    doc.core_properties.subject = "计算机软件著作权登记 · 文档鉴别材料"
    out = Path(args.out).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)

    pages = stats.pages()
    print(f"sections : {len(sections)}")
    print(f"est pages: {pages}  (字 {stats.chars} / 表行 {stats.table_rows} / 图 {stats.images})")
    print(f"images   : {images}")
    print(f"chars    : {chars}")
    print(f"out      : {out}")
    print("##DSH_JSON## " + json.dumps({
        "ok": True,
        "docx": str(out),
        "software_name": name,
        "version": version,
        "sections": len(sections),
        "images": images,
        "tables": tables,
        "chars": chars,
        "est_pages": pages,
        "thin": pages < 20,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
